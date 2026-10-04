#!/usr/bin/env python3
"""
Praxis content export.

Reads the authoring workbook and emits validated JSON for the site.
Refuses to write anything if validation fails, so a broken library
never reaches the repo.

    python3 export_library.py praxis-concept-library.xlsx -o ./content

Produces:
    content/concepts.json
    content/books.json
    content/index.json      (domains, identities, counts, cheap for the wheel)

Requires: openpyxl   (pip install openpyxl)
"""

import argparse, json, re, sys, unicodedata
from pathlib import Path
from collections import defaultdict

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is not installed.  Run:  pip install openpyxl")


# ---------------------------------------------------------------- helpers

def slug(text):
    """Stable url-safe id from a canonical name."""
    text = unicodedata.normalize("NFKD", str(text))
    text = text.encode("ascii", "ignore").decode()
    text = re.sub(r"[^\w\s-]", "", text).strip().lower()
    text = re.sub(r"[\s_-]+", "-", text)
    return re.sub(r"^(the|a|an)-", "", text)


def cell(v):
    return "" if v is None else str(v).strip()


def split_list(v):
    """Semicolon-separated cell -> clean list."""
    return [p.strip() for p in cell(v).replace("\n", ";").split(";") if p.strip()]


def parse_mechanism(raw):
    """
    Split the mechanism field into ordered chains for the Map It exercise.
    Each line becomes one chain; '->' separates its steps.
    A line with no arrow is kept as prose and flagged.
    """
    chains, prose = [], []
    for line in cell(raw).split("\n"):
        line = line.strip()
        if not line:
            continue
        if "->" in line or "→" in line:
            steps = [s.strip() for s in re.split(r"->|→", line) if s.strip()]
            if len(steps) >= 2:
                chains.append(steps)
            else:
                prose.append(line)
        else:
            prose.append(line)
    return chains, prose


class Report:
    def __init__(self):
        self.errors = defaultdict(list)
        self.warnings = defaultdict(list)

    def error(self, where, msg):
        self.errors[where].append(msg)

    def warn(self, where, msg):
        self.warnings[where].append(msg)

    def dump(self):
        for bucket, label in ((self.warnings, "WARNING"), (self.errors, "ERROR")):
            if not bucket:
                continue
            print()
            for where in sorted(bucket):
                print(f"{label}  {where}")
                for m in bucket[where]:
                    print(f"         - {m}")
        return not self.errors


# ---------------------------------------------------------------- reading

def read_sheet(ws):
    """Rows as dicts keyed by header, skipping blank rows."""
    headers = [cell(c.value) for c in ws[1]]
    out = []
    for row in ws.iter_rows(min_row=2):
        values = [c.value for c in row]
        if not any(cell(v) for v in values):
            continue
        out.append({h: values[i] if i < len(values) else None
                    for i, h in enumerate(headers) if h})
    return out


def read_vocabulary(ws):
    vocab = defaultdict(set)
    headers = [cell(c.value) for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        for i, h in enumerate(headers):
            if h and i < len(row):
                v = cell(row[i].value)
                if v:
                    vocab[h].add(v)
    return vocab


# ---------------------------------------------------------------- building

REQUIRED_ALL = ["One Sentence", "Definition", "Mechanism (ordered chain)",
                "Common Misconception", "Nearest Neighbour", "Discriminate Scenario"]
REQUIRED_INTEGRATION = ["Praxis Exercise", "Identity Tags"]


def build_books(raw_books, rep):
    books, seen = {}, set()
    for r in raw_books:
        bid = cell(r.get("Book ID"))
        title = cell(r.get("Title"))
        if not bid and not title:
            continue
        where = f"Books / {title or '(untitled)'}"
        if not bid:
            rep.error(where, "has a title but no Book ID")
            continue
        if not re.fullmatch(r"BK\d{3,}", bid):
            rep.error(where, f"Book ID '{bid}' should look like BK001")
        if bid in seen:
            rep.error(where, f"duplicate Book ID '{bid}'")
        seen.add(bid)
        books[bid] = {
            "id": bid,
            "title": title,
            "author": cell(r.get("Author")),
            "year": cell(r.get("Year")),
            "summaryStatus": cell(r.get("Summary status")),
            "oneLineTake": cell(r.get("One-line take")),
            "primaryDomain": cell(r.get("Primary Domain")),
            "concepts": [],
        }
    return books


def build_concepts(raw, vocab, books, rep):
    concepts, by_name = {}, {}

    for r in raw:
        name = cell(r.get("Canonical Name"))
        if not name:
            continue
        status = cell(r.get("Status"))
        if status != "Done":
            continue                      # only Done rows ship

        pid = cell(r.get("Praxis ID"))
        where = f"Concepts / {pid or '??'} {name}"
        if not pid:
            rep.error(where, "has no Praxis ID. Every concept needs its NN.NN id")
            continue
        cid = pid
        if cid in concepts:
            rep.error(where, f"duplicate Praxis ID '{cid}'")
            continue

        tier = cell(r.get("Tier")).lower()
        if tier not in ("discovery", "integration"):
            rep.error(where, f"Tier must be Discovery or Integration, got '{tier or 'blank'}'")

        # required fields
        missing = [f for f in REQUIRED_ALL if not cell(r.get(f))]
        if tier == "integration":
            missing += [f for f in REQUIRED_INTEGRATION if not cell(r.get(f))]
        for f in missing:
            rep.error(where, f"status is Done but '{f}' is empty")

        # a Discovery concept must NOT carry an exercise
        if tier == "discovery" and cell(r.get("Praxis Exercise")):
            rep.error(where, "Discovery tier must not have a Praxis Exercise. "
                             "Either clear it or promote the concept to Integration")

        if not cell(r.get("Evidence")):
            rep.error(where, "status is Done but 'Evidence' is empty. Every finished concept "
                             "has to say how well supported it is")

        if cell(r.get("Type")) in ("Mythic Concept", "Archetype") and not cell(r.get("Story")):
            rep.error(where, "a myth or archetype needs a Story. Without it the concept "
                             "reaches the reader as a definition, which is not what it is")

        # controlled vocabulary
        for field, key in (("Primary Domain", "Domains"), ("Type", "Types"),
                           ("Difficulty", "Difficulty"), ("Evidence", "Evidence")):
            v = cell(r.get(field))
            if v and v not in vocab.get(key, set()):
                rep.error(where, f"{field} '{v}' is not in the Vocabulary sheet")

        # identity tags
        tags = split_list(r.get("Identity Tags"))
        for t in tags:
            if t not in vocab.get("Identities", set()):
                rep.error(where, f"Identity Tag '{t}' is not in the Vocabulary sheet")
        if tier == "integration" and not tags:
            rep.error(where, "Integration tier needs at least one Identity Tag "
                             "or it can never be surfaced by a goal")
        if tier == "discovery" and not tags:
            rep.warn(where, "no Identity Tag, so this concept can appear on the wheel but "
                            "will never be weighted toward anyone's goal")

        # books
        bids = split_list(r.get("Book IDs"))
        for b in bids:
            if b not in books:
                rep.error(where, f"Book ID '{b}' does not exist on the Books sheet")
            else:
                books[b]["concepts"].append(cid)

        # mechanism
        chains, prose = parse_mechanism(r.get("Mechanism (ordered chain)"))
        if not chains and cell(r.get("Mechanism (ordered chain)")):
            rep.error(where, "Mechanism has no '->' arrows, so Map It cannot be generated. "
                             "Rewrite it as an ordered chain.")
        if prose and chains:
            rep.warn(where, "part of Mechanism has no arrows and was kept as prose")

        # scenario must not name the concept
        scenario = cell(r.get("Discriminate Scenario"))
        bare = re.sub(r"^(the|a|an)\s+", "", name, flags=re.I)
        if scenario and bare.lower() in scenario.lower():
            rep.error(where, f"Discriminate Scenario contains the concept name '{bare}'. "
                             "The Spot rung collapses if the answer is in the question.")
        words = len(scenario.split())
        if scenario and not (30 <= words <= 90):
            rep.warn(where, f"Discriminate Scenario is {words} words (aim 40-70)")

        concepts[cid] = {
            "id": cid,
            "slug": slug(name),
            "name": name,
            "tier": tier,
            "domain": cell(r.get("Primary Domain")),
            "type": cell(r.get("Type")),
            "legacyType": cell(r.get("Legacy Type")),
            "tradition": cell(r.get("Tradition")),
            "difficulty": cell(r.get("Difficulty")).lower(),
            "evidence": cell(r.get("Evidence")),
            "origin": cell(r.get("Origin")),
            "date": cell(r.get("Date")),
            "oneSentence": cell(r.get("One Sentence")),
            "story": cell(r.get("Story")),
            "definition": cell(r.get("Definition")),
            "mechanism": {"chains": chains, "prose": prose},
            "misconception": cell(r.get("Common Misconception")),
            "nearestNeighbour": cell(r.get("Nearest Neighbour")),
            "discriminateScenario": scenario,
            "praxisExercise": cell(r.get("Praxis Exercise")),
            "identityTags": tags,
            "related": [],                      # filled below
            "books": bids,
            "author": cell(r.get("Author")),
        }
        by_name[name.lower()] = cid
        by_name[cid] = cid
        by_name[slug(name)] = cid

    # ---- related concepts, mirrored both ways
    edges = defaultdict(set)
    for r in raw:
        name = cell(r.get("Canonical Name"))
        if cell(r.get("Status")) != "Done" or not name:
            continue
        cid = cell(r.get("Praxis ID"))
        if cid not in concepts:
            continue
        for other in split_list(r.get("Related Concepts")):
            oid = by_name.get(other) or by_name.get(other.lower()) or by_name.get(slug(other))
            if not oid:
                rep.warn(f"Concepts / {name}",
                         f"Related Concept '{other}' is not a Done concept yet. "
                         "Edge skipped, it will appear once that row is finished")
                continue
            if oid == cid:
                continue
            edges[cid].add(oid)
            edges[oid].add(cid)          # mirror

    for cid, linked in edges.items():
        if cid in concepts:
            concepts[cid]["related"] = sorted(linked)

    return concepts


def build_index(concepts, books, vocab):
    by_identity, by_domain = defaultdict(list), defaultdict(list)
    for c in concepts.values():
        for t in c["identityTags"]:
            by_identity[t].append(c["id"])
        if c["domain"]:
            by_domain[c["domain"]].append(c["id"])
    return {
        "counts": {
            "concepts": len(concepts),
            "integration": sum(1 for c in concepts.values() if c["tier"] == "integration"),
            "discovery": sum(1 for c in concepts.values() if c["tier"] == "discovery"),
            "books": len(books),
        },
        "identities": {k: sorted(v) for k, v in sorted(by_identity.items())},
        "domains": {k: sorted(v) for k, v in sorted(by_domain.items())},
        "scenarioPool": sorted(c["id"] for c in concepts.values() if c["discriminateScenario"]),
        "vocabulary": {k: sorted(v) for k, v in vocab.items()},
    }


# ---------------------------------------------------------------- spinner

def update_spinner(wb, site_html, rep):
    """Rewrite the Discover spinner's inline DATA from the workbook.

    Discover shows the whole library, stubs included, because discovery is
    browsing. Comprehend shows only finished concepts, because nobody can be
    questioned on something nobody has written. One source, two audiences.

    Existing entries keep their book references exactly: those were authored
    on the page and are not in the workbook.
    """
    if not site_html.exists():
        rep.warn("Spinner", f"{site_html} not found, leaving the spinner alone")
        return None

    html = site_html.read_text(encoding="utf-8")
    m = re.search(r"const DATA = (\[.*?\]);", html, re.S)
    if not m:
        rep.warn("Spinner", "could not find the DATA array in the page")
        return None
    old = {d["id"]: d for d in json.loads(m.group(1))}

    books = {}
    for row in read_sheet(wb["Books"]):
        bid = cell(row.get("Book ID"))
        if bid:
            books[bid] = {"t": cell(row.get("Title")), "a": cell(row.get("Author")),
                          "y": cell(row.get("Year")), "k": "book"}

    out, added = [], 0
    for r in read_sheet(wb["Concepts"]):
        pid, name = cell(r.get("Praxis ID")), cell(r.get("Canonical Name"))
        if not pid or not name:
            continue
        prev = old.get(pid, {})
        refs = prev.get("r")
        if refs is None:
            refs = [books[b] for b in split_list(r.get("Book IDs")) if b in books]
        if pid not in old:
            added += 1
        out.append({
            "id": pid,
            "c": name,
            "d": cell(r.get("Primary Domain")),
            "t": cell(r.get("Legacy Type")) or cell(r.get("Type")),
            "s": cell(r.get("Tradition")),
            "b": prev.get("b"),
            "r": refs,
        })
    out.sort(key=lambda d: d["id"])

    new_html = html[:m.start(1)] + json.dumps(out, ensure_ascii=False) + html[m.end(1):]
    site_html.write_text(new_html, encoding="utf-8")
    return len(out), added


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description="Export the Praxis concept library to JSON.")
    ap.add_argument("workbook", help="the authoring .xlsx")
    ap.add_argument("-o", "--out", default="./content", help="output directory (default ./content)")
    ap.add_argument("--allow-warnings", action="store_true",
                    help="(default) warnings never block; kept for symmetry")
    args = ap.parse_args()

    path = Path(args.workbook)
    if not path.exists():
        sys.exit(f"No such file: {path}")

    wb = load_workbook(path, data_only=True)
    for required in ("Concepts", "Books", "Vocabulary"):
        if required not in wb.sheetnames:
            sys.exit(f"Workbook is missing the '{required}' sheet.")

    rep = Report()
    vocab = read_vocabulary(wb["Vocabulary"])
    books = build_books(read_sheet(wb["Books"]), rep)
    concepts = build_concepts(read_sheet(wb["Concepts"]), vocab, books, rep)

    for b in books.values():
        b["concepts"] = sorted(set(b["concepts"]))

    print(f"\nRead {len(concepts)} finished concepts and {len(books)} books.")
    ok = rep.dump()

    if not ok:
        n = sum(len(v) for v in rep.errors.values())
        print(f"\n{n} error(s). Nothing was written. Fix the workbook and run again.\n")
        sys.exit(1)

    if not concepts:
        print("\nNo concepts are marked 'Done' yet, so this writes an empty library.")
        print("This is expected while the 300 rows are still stubs.\n")

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    payload = [
        ("concepts.json", sorted(concepts.values(), key=lambda c: c["id"])),
        ("books.json", sorted(books.values(), key=lambda b: b["id"])),
        ("index.json", build_index(concepts, books, vocab)),
    ]
    site_html = Path(args.workbook).resolve().parent.parent / "site" / "index.html"
    spin = update_spinner(wb, site_html, rep)

    for fname, data in payload:
        (out / fname).write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n",
                                 encoding="utf-8")
        print(f"  wrote {out / fname}")

    if spin:
        print(f"  spinner now carries {spin[0]} concepts ({spin[1]} newly added)")

    nwarn = sum(len(v) for v in rep.warnings.values())
    print(f"\nDone{f', with {nwarn} warning(s) above, nothing blocking' if nwarn else ''}.\n")


if __name__ == "__main__":
    main()
